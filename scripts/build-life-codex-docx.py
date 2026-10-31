#!/usr/bin/env python3
"""Build the Arabic AURA Word codex from the exact runtime quest generator output."""
import json
import os
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
OUTPUT = ROOT / "public" / "downloads" / "AURA-Life-Quests-22500-Ar.docx"
quests = json.loads((ARTIFACTS / "life-quests-ar.json").read_text(encoding="utf-8"))
manifest = json.loads((ARTIFACTS / "life-codex-manifest.json").read_text(encoding="utf-8"))
assert len(quests) == 22500 == manifest["total"]

skill_names = {
    "discipline": "الانضباط", "iron-body": "الجسد الحديدي", "deep-work": "العمل العميق",
    "golden-touch": "اللمسة الذهبية", "second-wind": "النَفَس الثاني", "inner-peace": "السلام الداخلي",
    "swift-strike": "الضربة السريعة", "mind-palace": "قصر الذهن",
}
rank_names = {
    "F": "المبتدئ", "E": "المستكشف", "D": "المتدرب", "C": "الممارس", "B": "المتقدم",
    "A": "المتقن", "S": "الخبير", "SS": "البطل", "SSS": "الأسطوري",
}
rank_colors = {rank["rank"]: rank["color"].lstrip("#") for rank in manifest["ranks"]}


def rtl_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = alignment
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        p_pr.append(bidi)
    return paragraph


def rtl_run(run, size=None, bold=None, color=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:bidi"), "ar-MA")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade_paragraph(paragraph, fill="F7F8FA", border="D9DDE5"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    right = OxmlElement("w:right")
    right.set(qn("w:val"), "single")
    right.set(qn("w:sz"), "18")
    right.set(qn("w:space"), "7")
    right.set(qn("w:color"), border)
    borders.append(right)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "3")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "E5E7EB")
    borders.append(bottom)
    p_pr.append(borders)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), "70")
    spacing.set(qn("w:after"), "90")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:right"), "130")
    ind.set(qn("w:left"), "100")
    p_pr.append(ind)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    return run


def add_heading(document, text, level, color=None):
    paragraph = document.add_heading(level=level)
    rtl_paragraph(paragraph)
    run = paragraph.add_run(text)
    rtl_run(run, 22 if level == 1 else 16, True, color or ("A87528" if level == 1 else "263449"))
    return paragraph


def add_body(document, text, bold=False, color="374151", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph = document.add_paragraph()
    rtl_paragraph(paragraph, align)
    rtl_run(paragraph.add_run(text), size, bold, color)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    return paragraph


def set_cell(cell, text, bold=False, fill=None, color="263449"):
    if fill:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    rtl_paragraph(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    rtl_run(paragraph.add_run(str(text)), 9, bold, color)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.45)
section.right_margin = Cm(1.45)
section.header_distance = Cm(0.65)
section.footer_distance = Cm(0.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10)
normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
normal.paragraph_format.space_after = Pt(4)
for style_name, size, color in (("Title", 30, "111319"), ("Heading 1", 22, "A87528"), ("Heading 2", 16, "263449"), ("Heading 3", 12, "263449")):
    style = styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(8)
quest_style = styles.add_style("AURA Quest Card", WD_STYLE_TYPE.PARAGRAPH)
quest_style.base_style = normal
quest_style.font.name = "Arial"
quest_style.font.size = Pt(9)
quest_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
quest_style.paragraph_format.space_after = Pt(4)
quest_style.paragraph_format.keep_together = True

header = section.header.paragraphs[0]
rtl_paragraph(header)
rtl_run(header.add_run("AURA  /  موسوعة بطاقات المهام الحياتية  /  22,500 مهمة"), 8, True, "8B6A35")
footer = section.footer.paragraphs[0]
rtl_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
rtl_run(footer.add_run("AURA  •  "), 8, False, "6B7280")
rtl_run(add_field(footer, "PAGE"), 8, False, "6B7280")

# Cover
p = doc.add_paragraph()
rtl_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.space_before = Pt(95)
rtl_run(p.add_run("AURA"), 42, True, "B58332")
p = doc.add_paragraph()
rtl_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
rtl_run(p.add_run("موسوعة بطاقات المهام الحياتية"), 27, True, "111319")
p = doc.add_paragraph()
rtl_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
rtl_run(p.add_run("22,500 مهمة مختلفة  •  9 رتب  •  6 مسارات  •  4 لغات داخل التطبيق"), 14, True, "4B5563")
p.paragraph_format.space_before = Pt(16)
add_body(doc, "النسخة العربية المرجعية لمحتوى AURA 1.11 — مولّدة من المصدر نفسه المستخدم في Deck of Fate وQuest Codex.", color="6B7280", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
add_body(doc, "نسخة تفصيلية: اسم ووصف وطريقة إنجاز ومكافآت ومدة مستقلة لكل مهمة.", bold=True, color="4B5563", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
add_body(doc, "F  →  E  →  D  →  C  →  B  →  A  →  S  →  SS  →  SSS", bold=True, color="A87528", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
add_body(doc, "31 أكتوبر 2026", color="6B7280", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

add_heading(doc, "منهجية بناء الموسوعة", 1)
add_body(doc, "تتكون كل رتبة من 2,500 مهمة: 500 دينية، و400 نفسية، و400 جسدية، و400 عقلية، و400 لتطوير الذات، و400 للعلاقات الاجتماعية. تتدرج المهام بالجودة والثبات والتعقيد والمراجعة، لا بالإفراط وحده. لكل بطاقة معرّف ثابت وصورة مولّدة ونص ومرجع إرشادي ومهارات مرتبطة داخل التطبيق. يعرض هذا الملف داخل كل بطاقة اسم المهمة، ووصفها وتفاصيل إنجازها في سطرين مستقلين على الأقل، ثم نقاط الخبرة والعملات الذهبية والماس والمدة المقدرة بدقة.")
add_body(doc, "المحتوى توليف تطبيقي من المراجع، وليس اقتباسًا حرفيًا من كل كتاب ولا فتوى ولا وصفة علاجية. القرآن والسنة الصحيحة هما المرجع الديني الأول؛ وكتب التفسير والفقه والسيرة والتربية تُقرأ بسياقها وبالرجوع إلى أهل العلم. ومراجع الطب النفسي وDSM-5-TR وICD-11 للفهم المهني والإحالة، وليست لإنشاء تشخيص ذاتي أو تغيير دواء.")
add_body(doc, "التأثير في قرارات الجنس الآخر يعني هنا: التواصل الصادق، والأسئلة المفتوحة، وعرض الدليل والمصلحة المشتركة، وحرية الرفض، واحترام الرضا والخصوصية والحدود الشرعية والمهنية. لا تتضمن الموسوعة تكتيكات تلاعب أو ضغط أو غيرة مقصودة أو استغلال ضعف أو ملاحقة بعد الرفض.")
add_body(doc, "تنبيه صحي: أوقف النشاط عند ألم حاد أو دوخة أو ازدياد انزعاج نفسي. اطلب مختصًا مؤهلًا عند الحاجة، ولا توقف علاجًا أو دواءً بسبب بطاقة. عند وجود خطر فوري على النفس أو الغير، اطلب خدمات الطوارئ المحلية أو شخصًا موثوقًا فورًا؛ المهام ليست خطة أزمة بديلة.")

add_heading(doc, "ظهور البطاقات حسب المستوى", 1)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, ("مستوى اللاعب", "الرتب المتاحة", "ملاحظة")):
    set_cell(cell, text, True, "E8DCC5")
for rank in manifest["ranks"]:
    cells = table.add_row().cells
    index = manifest["ranks"].index(rank)
    available = " + ".join(item["rank"] for item in manifest["ranks"][:index + 1])
    set_cell(cells[0], rank["levelRange"])
    set_cell(cells[1], available, True, rank_colors[rank["rank"]])
    set_cell(cells[2], "تُضاف الرتبة الجديدة وتبقى جميع الرتب السابقة")

add_heading(doc, "التوزيع الكامل", 1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, ("المسار", "لكل رتبة", "المجموع عبر 9 رتب", "المعرّف")):
    set_cell(cell, text, True, "E8DCC5")
for category in manifest["categories"]:
    cells = table.add_row().cells
    set_cell(cells[0], category["label"])
    set_cell(cells[1], f"{category['perRank']:,}")
    set_cell(cells[2], f"{category['total']:,}")
    set_cell(cells[3], category["prefix"])

add_heading(doc, "فهرس قابل للتحديث", 1)
p = doc.add_paragraph()
rtl_paragraph(p)
add_field(p, 'TOC \\o "1-2" \\h \\z \\u')
add_body(doc, "في Microsoft Word: حدّد الفهرس ثم اختر Update Field لتحديث أرقام الصفحات.", color="6B7280", size=9)
doc.add_page_break()

add_heading(doc, "المراجع ومنهج الاستفادة منها", 1)
for code, title, note in manifest["references"]:
    p = doc.add_paragraph(style="AURA Quest Card")
    rtl_paragraph(p)
    rtl_run(p.add_run(f"[{code}] {title}"), 9, True, "8B6A35")
    rtl_run(p.add_run(f"\n{note}"), 9, False, "4B5563")
    shade_paragraph(p, "F8F6F1", "C7A76A")
doc.add_page_break()

by_section = defaultdict(list)
for quest in quests:
    by_section[(quest["category"], quest["rank"])].append(quest)

category_map = {category["id"]: category for category in manifest["categories"]}
for category_index, category in enumerate(manifest["categories"]):
    add_heading(doc, f"{category_index + 1}. {category['label']}", 1)
    add_body(doc, f"المجموع: {category['total']:,} مهمة. في كل رتبة: {category['perRank']:,} مهمة. رمز المسار: {category['prefix']}.")
    for rank in manifest["ranks"]:
        rank_id = rank["rank"]
        section_quests = by_section[(category["id"], rank_id)]
        assert len(section_quests) == category["perRank"]
        add_heading(doc, f"الرتبة {rank_id} — {rank_names[rank_id]}  |  {len(section_quests):,} مهمة  |  مستوى {rank['levelRange']}", 2, rank_colors[rank_id])
        add_body(doc, "تتضمن كل بطاقة في هذا القسم وصفًا وتفاصيل إنجاز، ثم نقاط الخبرة والعملات الذهبية والماس والمدة المقدرة الخاصة بالمهمة.", color="6B7280", size=9)
        for quest in section_quests:
            p = doc.add_paragraph(style="AURA Quest Card")
            rtl_paragraph(p)
            rtl_run(p.add_run(quest["id"]), 8, True, rank_colors[rank_id])
            rtl_run(p.add_run(f"  |  الرتبة {rank_id}  |  اسم المهمة: {quest['baseAction']}"), 8, True, "4B5563")
            rtl_run(p.add_run(f"\nالوصف: {quest['text']}"), 9, True, "17202A")
            rtl_run(p.add_run(f"\nالتفاصيل العملية ودليل الإنجاز: {quest['completion']}"), 8, False, "4B5563")
            rewards = f"نقاط الخبرة: {quest['xp']:,} XP  |  العملات الذهبية: {quest['gold']:,}  |  الماس: {quest['diamonds']:,}"
            rtl_run(p.add_run(f"\nالمكافآت عند الإنجاز: {rewards}"), 8, True, "8B6A35")
            skills = "، ".join(skill_names.get(skill, skill) for skill in quest["skillIds"])
            rtl_run(p.add_run(f"\nالمدة المقدرة: {quest['duration']} دقيقة  |  المهارات: {skills}  |  المراجع: {' · '.join(quest['sourceCodes'])}"), 8, False, "5F6672")
            shade_paragraph(p, "F8F9FB", rank_colors[rank_id])
        doc.add_page_break()

# Ask Word to refresh fields (TOC/page numbers) when opening.
settings = doc.settings._element
update_fields = OxmlElement("w:updateFields")
update_fields.set(qn("w:val"), "true")
settings.append(update_fields)
doc.core_properties.title = "AURA — موسوعة بطاقات المهام الحياتية: 22,500 مهمة"
doc.core_properties.subject = "مهام دينية ونفسية وجسدية وعقلية وتطوير ذات وعلاقات، من F إلى SSS"
doc.core_properties.author = "AURA"
doc.core_properties.keywords = "AURA, مهام, تطوير الذات, صحة, علاقات, دين, RPG"
doc.core_properties.comments = "Generated from the same runtime content used by AURA 1.11."
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(f"Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes) with {len(quests):,} tasks.")
